import io
import re

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Image as PDFImage
)

from PIL import Image


def clean_text(text):

    replacements = {
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2013": "-",
        "\u2014": "-",
        "\u2026": "..."
    }

    for old, new in replacements.items():
        text = text.replace(old, new)

    return text.strip()


def create_txt(text):

    text = clean_text(text)

    return (
        text + "\n"
    ).encode("utf-8")


def is_heading(line):

    line = line.strip()

    if not line:
        return False

    if re.match(
        r"^(\d+[\.\)]|SECTION|ARTICLE|CLAUSE)",
        line,
        re.IGNORECASE
    ):
        return True

    if line.isupper() and len(line) < 100:
        return True

    return False


def create_docx(
    text,
    document_type,
    logo=None
):

    document = Document()

    section = document.sections[0]

    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    # Normal style

    normal = document.styles["Normal"]

    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)

    # Logo

    if logo:

        try:

            paragraph = document.add_paragraph()

            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
            )

            run = paragraph.add_run()

            run.add_picture(
                io.BytesIO(logo),
                width=Inches(1.5)
            )

        except Exception:
            pass

    # Title

    title = document.add_paragraph()

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    title_run = title.add_run(
        document_type.upper()
    )

    title_run.bold = True
    title_run.font.name = "Times New Roman"
    title_run.font.size = Pt(16)

    document.add_paragraph()

    # Body

    lines = clean_text(text).splitlines()

    for line in lines:

        line = line.strip()

        if not line:
            continue

        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_after = Pt(7)

        if is_heading(line):

            run = paragraph.add_run(line)

            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)

        else:

            run = paragraph.add_run(line)

            run.font.name = "Times New Roman"
            run.font.size = Pt(11)

    # Footer

    footer = section.footer

    footer_paragraph = footer.paragraphs[0]

    footer_paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    footer_run = footer_paragraph.add_run(
        "LegalEase - AI-generated legal draft. "
        "Please review with a qualified legal professional."
    )

    footer_run.font.size = Pt(8)

    # Save to memory

    output = io.BytesIO()

    document.save(output)

    return output.getvalue()


def create_pdf(
    text,
    document_type,
    logo=None
):

    output = io.BytesIO()

    pdf = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=20 * mm,
        leftMargin=20 * mm,
        topMargin=25 * mm,
        bottomMargin=25 * mm
    )

    styles = getSampleStyleSheet()

    title_style = styles["Title"]

    title_style.fontName = "Helvetica-Bold"
    title_style.fontSize = 16
    title_style.leading = 20
    title_style.alignment = 1

    heading_style = styles["Heading2"]

    heading_style.fontName = "Helvetica-Bold"
    heading_style.fontSize = 11

    body_style = styles["BodyText"]

    body_style.fontName = "Helvetica"
    body_style.fontSize = 9.5
    body_style.leading = 14

    elements = []

    # Logo

    if logo:

        try:

            logo_buffer = io.BytesIO(logo)

            image = PDFImage(
                logo_buffer,
                width=40 * mm,
                height=20 * mm
            )

            image.hAlign = "CENTER"

            elements.append(image)

            elements.append(
                Spacer(1, 8)
            )

        except Exception:
            pass

    # Title

    elements.append(
        Paragraph(
            document_type.upper(),
            title_style
        )
    )

    elements.append(
        Spacer(1, 15)
    )

    # Document content

    for line in clean_text(text).splitlines():

        line = line.strip()

        if not line:
            elements.append(
                Spacer(1, 5)
            )

            continue

        escaped = (
            line
            .replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )

        if is_heading(line):

            elements.append(
                Paragraph(
                    escaped,
                    heading_style
                )
            )

        else:

            elements.append(
                Paragraph(
                    escaped,
                    body_style
                )
            )

        elements.append(
            Spacer(1, 4)
        )

    def footer(canvas, doc):

        canvas.saveState()

        canvas.setFont(
            "Helvetica",
            7
        )

        canvas.drawCentredString(
            A4[0] / 2,
            10 * mm,
            "LegalEase - AI-generated legal draft. "
            "Review before use."
        )

        canvas.drawRightString(
            A4[0] - 15 * mm,
            10 * mm,
            f"Page {doc.page}"
        )

        canvas.restoreState()

    pdf.build(
        elements,
        onFirstPage=footer,
        onLaterPages=footer
    )

    return output.getvalue()